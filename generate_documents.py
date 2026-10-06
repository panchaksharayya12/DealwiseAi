import os
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

from pptx import Presentation
from pptx.util import Inches as PptxInches, Pt as PptxPt
from pptx.dml.color import RGBColor as PptxRGBColor
from pptx.enum.text import PP_ALIGN

# ==========================================
# 1. GENERATE WORD DOCUMENT (.docx)
# ==========================================

def create_word_report():
    doc = Document()

    # Page Margins
    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    # Styles
    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Calibri'
    normal_style.font.size = Pt(11)
    normal_style.font.color.rgb = RGBColor(30, 30, 35)

    # Helper function for cell background shading
    def set_cell_background(cell, fill_hex):
        shading = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
        cell._tc.get_or_add_tcPr().append(shading)

    # Header / Title
    title_p = doc.add_paragraph()
    title_run = title_p.add_run("DealWise AI")
    title_run.font.size = Pt(28)
    title_run.font.bold = True
    title_run.font.color.rgb = RGBColor(10, 10, 15)

    subtitle_p = doc.add_paragraph()
    sub_run = subtitle_p.add_run("Comprehensive Technical & Product Diligence Report\nTagline: \"Know the deal before you buy.\"")
    sub_run.font.size = Pt(13)
    sub_run.font.italic = True
    sub_run.font.color.rgb = RGBColor(90, 90, 100)

    # Metadata Block
    meta_p = doc.add_paragraph()
    meta_p.add_run("Author: Panchaksharayya\n").bold = True
    meta_p.add_run("Repository: https://github.com/panchaksharayya12/DealwiseAi\n")
    meta_p.add_run("Date: October 2026 | Version: 1.0.0 (Production Release)\n")
    meta_p.paragraph_format.space_after = Pt(18)

    def add_heading_1(text):
        h = doc.add_heading(text, level=1)
        h.runs[0].font.size = Pt(16)
        h.runs[0].font.bold = True
        h.runs[0].font.color.rgb = RGBColor(15, 23, 42)
        h.paragraph_format.space_before = Pt(14)
        h.paragraph_format.space_after = Pt(6)
        return h

    def add_heading_2(text):
        h = doc.add_heading(text, level=2)
        h.runs[0].font.size = Pt(13)
        h.runs[0].font.bold = True
        h.runs[0].font.color.rgb = RGBColor(51, 65, 85)
        h.paragraph_format.space_before = Pt(10)
        h.paragraph_format.space_after = Pt(4)
        return h

    # Section 1: Executive Summary
    add_heading_1("1. Executive Summary")
    doc.add_paragraph(
        "Real estate purchases represent significant financial commitments for retail buyers, institutional investors, "
        "and syndicates. However, the majority of property evaluations rely on promotional broker marketing, surface-level "
        "asking prices, or subjective visual impressions rather than quantitative scrutiny."
    )
    doc.add_paragraph(
        "DealWise AI is a full-stack real-estate deal analyzer engineered to answer the critical question: "
        "\"Is this property actually a good deal?\" By coupling 100% deterministic actuarial and financial calculations "
        "with an objective 5-pillar scoring algorithm and contextual AI synthesis, DealWise AI equips buyers with data-driven "
        "clarity, risk warnings, multi-property benchmarking, and verifiable investment diligence reports."
    )

    # Section 2: Problem Statement & Market Opportunity
    add_heading_1("2. Problem Statement & Market Opportunity")
    add_heading_2("2.1 The Value vs. Price Distortion")
    doc.add_paragraph(
        "Homebuyers frequently conflate the asking price of a property with its investment value. Without calculating capitalization "
        "rates, net operating yield, or total financing costs, buyers regularly commit capital to properties that deliver below-inflation "
        "returns."
    )
    add_heading_2("2.2 Negative Carry and Debt Traps")
    doc.add_paragraph(
        "High Loan-to-Value (LTV) loans paired with recurring society maintenance fees often create an uncalculated monthly cash deficit. "
        "When the combined monthly EMI and maintenance charges exceed the gross rent, owners face substantial recurring out-of-pocket expenses."
    )
    add_heading_2("2.3 Artificial Intelligence Hallucinations in Financial Modeling")
    doc.add_paragraph(
        "General-purpose Large Language Models (LLMs) frequently hallucinate amortization schedules, compound interest tallies, "
        "and yields. DealWise AI solves this by keeping all mathematics strictly within deterministic TypeScript code routines, "
        "reserving AI exclusively for qualitative interpretation and due-diligence guidance."
    )

    # Section 3: System Architecture
    add_heading_1("3. System Architecture & Tech Stack")
    doc.add_paragraph(
        "DealWise AI is structured with a decoupled, high-throughput architecture separating the client interface, REST API services, "
        "deterministic calculation utilities, and optional cloud persistence."
    )

    # Architecture Table
    arch_table = doc.add_table(rows=1, cols=3)
    arch_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr = arch_table.rows[0].cells
    hdr[0].text = "Tier"
    hdr[1].text = "Technology"
    hdr[2].text = "Functional Purpose"
    for c in hdr:
        set_cell_background(c, "F1F5F9")
        c.paragraphs[0].runs[0].font.bold = True

    arch_rows = [
        ("Frontend Client", "React 18, TypeScript, Vite", "Interactive UI, input validation, state management, and real-time counter rendering."),
        ("Design System", "Tailwind CSS, Liquid Glass", "Minimal VEX-inspired aesthetic: pure black (#000000), white, zinc gray, specular borders, zero gradients."),
        ("Visual Analytics", "Recharts (SVG)", "Responsive 5-year capital growth and accumulated rental income projections with interactive appreciation slider."),
        ("Document Engine", "jsPDF", "Client-side vector generation of formal executive due-diligence reports."),
        ("Backend Server", "Express, Node.js, tsx", "REST API validating property inputs, computing deterministic metrics, and serving multi-property comparison logic."),
        ("AI Layer", "OpenAI GPT-4o-mini", "Context-aware qualitative synthesis, seller negotiation strategies, and conversational queries with automatic local fallback."),
        ("Data Persistence", "Supabase & localStorage", "Dual-mode persistence: cloud PostgreSQL synchronization when configured, and automatic browser storage fallback.")
    ]

    for tier, tech, purpose in arch_rows:
        row = arch_table.add_row().cells
        row[0].text = tier
        row[1].text = tech
        row[2].text = purpose
        row[0].paragraphs[0].runs[0].font.bold = True
        for c in row:
            c.paragraphs[0].runs[0].font.size = Pt(9.5)

    doc.add_paragraph().paragraph_format.space_after = Pt(10)

    # Section 4: Deterministic Financial Modeling
    add_heading_1("4. Deterministic Financial Modeling & Actuarial Formulas")
    doc.add_paragraph(
        "All calculations executed by DealWise AI are deterministic, auditable, and mathematically grounded. "
        "The following core formulas govern the evaluation:"
    )

    formulas = [
        ("Price per Square Foot", "Price / sq.ft = Asking Price / Built-up Area"),
        ("Gross Rental Yield", "Gross Yield (%) = (Expected Monthly Rent * 12 / Asking Price) * 100"),
        ("Net Rental Yield", "Net Yield (%) = ((Monthly Rent - Monthly Maintenance) * 12 / Asking Price) * 100"),
        ("Loan Amount", "Loan Amount = max(0, Asking Price - Down Payment)"),
        ("Monthly EMI (Standard Amortization)", "EMI = P * r * (1+r)^n / ((1+r)^n - 1), where P = Loan Amount, r = Annual Rate / (12 * 100), n = Tenure * 12"),
        ("Total Interest Payable", "Total Interest = (EMI * n) - Loan Amount"),
        ("Monthly Cashflow Carry", "Monthly Cashflow = Expected Monthly Rent - Monthly Maintenance - Monthly EMI"),
        ("Loan-to-Value (LTV) Ratio", "LTV (%) = (Loan Amount / Asking Price) * 100"),
        ("Cash-on-Cash ROI", "ROI (%) = (Annual Net Cashflow / Down Payment) * 100")
    ]

    for title, formula in formulas:
        p = doc.add_paragraph()
        p.add_run(f"• {title}: ").bold = True
        p.add_run(formula)
        p.paragraph_format.space_after = Pt(3)

    # Section 5: The 5-Pillar Deal Score Matrix
    add_heading_1("5. The Proprietary 5-Pillar Deal Score (0–100)")
    doc.add_paragraph(
        "DealWise AI aggregates five weighted quantitative pillars to establish a single, transparent Deal Score between 0 and 100:"
    )
    doc.add_paragraph(
        "Deal Score = (0.35 * Financial Health) + (0.20 * Rental Yield) + (0.20 * Price Efficiency) + (0.15 * Loan Burden) + (0.10 * Risk Buffer)"
    ).bold = True

    score_table = doc.add_table(rows=1, cols=3)
    score_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    s_hdr = score_table.rows[0].cells
    s_hdr[0].text = "Pillar"
    s_hdr[1].text = "Weight"
    s_hdr[2].text = "Quantitative Criteria"
    for c in s_hdr:
        set_cell_background(c, "F1F5F9")
        c.paragraphs[0].runs[0].font.bold = True

    score_rows = [
        ("Financial Health", "35%", "Evaluates net rental yield relative to cost of capital, monthly cashflow coverage, and cash-on-cash equity yield."),
        ("Rental Yield Power", "20%", "Benchmarks gross rental yield against metropolitan residential benchmarks (>5.5% = exceptional, 4.0-5.4% = good, <2.5% = poor)."),
        ("Price Efficiency", "20%", "Affordability and capitalization rate per square foot relative to the asset configuration and total capital outlay."),
        ("Loan Burden", "15%", "Assesses leverage stress through Loan-to-Value (LTV) percentage and monthly debt service coverage (EMI vs. expected rent)."),
        ("Risk Buffer", "10%", "Structural age depreciation liability, maintenance drag percentage, and dedicated parking inclusion.")
    ]

    for p_name, w_val, c_crit in score_rows:
        row = score_table.add_row().cells
        row[0].text = p_name
        row[1].text = w_val
        row[2].text = c_crit
        row[0].paragraphs[0].runs[0].font.bold = True
        for c in row:
            c.paragraphs[0].runs[0].font.size = Pt(9.5)

    doc.add_paragraph().paragraph_format.space_after = Pt(8)

    add_heading_2("5.1 Classification Thresholds")
    doc.add_paragraph("• 80 – 100: Strong Deal (Solid fundamentals, favorable yields, sustainable leverage, positive/balanced carry).")
    doc.add_paragraph("• 60 – 79: Fair Deal (Viable property with reasonable returns; minor negotiation or down payment increase advised).")
    doc.add_paragraph("• 40 – 59: Needs Review (Friction identified: compressed rental yield, elevated financing cost, or maintenance drag).")
    doc.add_paragraph("• 0 – 39: Risky Deal (Severe negative carry, excessive debt burden, or disproportionate asking price versus rental income).")

    # Section 6: Risk Taxonomy
    add_heading_1("6. Risk Diagnostics & Pre-Purchase Due-Diligence Checklist")
    doc.add_paragraph(
        "DealWise AI classifies operational, structural, and debt liabilities into High, Medium, and Low severity flags. "
        "Furthermore, each evaluation provides an 8-point due-diligence checklist that buyers must physically verify:"
    )

    checklist_items = [
        "Marketable Title Deed: Verify a 30-year unbroken chain of title documents and registered conveyance deeds.",
        "RERA Registration & Sanctions: Confirm state Real Estate Regulatory Authority approval, sanctioned plans, and Commencement/Occupancy Certificates (CC/OC).",
        "Encumbrance Certificate (EC): Obtain an up-to-date EC from the sub-registrar office covering at least 15–30 years to verify absence of mortgages.",
        "Local Micro-Market Comparables: Cross-reference actual registered sale deed prices of comparable units within a 500-meter radius.",
        "Society Maintenance Health: Inspect association balance sheets, sinking fund reserves, and pending special repair assessments.",
        "Rental Absorption Rates: Confirm prevailing rental demand and tenant turnover rates directly with local independent brokers.",
        "Municipal Clearance: Confirm payment receipts for property tax, electricity, water, and sewerage charges.",
        "Bank Project Approval: Ensure the development is pre-approved for mortgage financing (APF) by major financial institutions."
    ]

    for idx, item in enumerate(checklist_items, 1):
        doc.add_paragraph(f"{idx}. {item}")

    # Section 7: REST API Endpoints
    add_heading_1("7. REST API Endpoints Specification")
    api_table = doc.add_table(rows=1, cols=4)
    api_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    a_hdr = api_table.rows[0].cells
    a_hdr[0].text = "Method"
    a_hdr[1].text = "Endpoint"
    a_hdr[2].text = "Payload"
    a_hdr[3].text = "Description"
    for c in a_hdr:
        set_cell_background(c, "F1F5F9")
        c.paragraphs[0].runs[0].font.bold = True

    api_rows = [
        ("GET", "/api/health", "None", "Service uptime, timestamp, and active integration flags."),
        ("POST", "/api/analyze", "{ property, appreciationRate }", "Validates input, computes metrics, generates score, and returns analysis."),
        ("POST", "/api/ai/chat", "{ property, metrics, score, userQuestion }", "Conversational AI responses grounded in the current property context."),
        ("POST", "/api/compare", "{ properties: PropertyInput[] }", "Multi-property benchmarking with category winner determinations."),
        ("GET", "/api/analyses", "None", "Retrieves stored property evaluations from Supabase or memory store."),
        ("POST", "/api/analyses", "SavedAnalysisRecord", "Persists evaluated deal records."),
        ("DELETE", "/api/analyses/:id", "URL param: id", "Deletes a stored property analysis."),
        ("POST", "/api/report", "{ analysis }", "Validates and formats metadata for investment report exports.")
    ]

    for m, ep, pl, ds in api_rows:
        row = api_table.add_row().cells
        row[0].text = m
        row[1].text = ep
        row[2].text = pl
        row[3].text = ds
        row[0].paragraphs[0].runs[0].font.bold = True
        for c in row:
            c.paragraphs[0].runs[0].font.size = Pt(9)

    doc.add_paragraph().paragraph_format.space_after = Pt(10)

    # Section 8: Quality Assurance & Verification
    add_heading_1("8. Verification & Quality Assurance Summary")
    doc.add_paragraph(
        "The system has completed rigorous verification testing across environments: "
        "TypeScript static type checks pass with zero errors (tsc -b), the Vite production build minifies without warnings, "
        "universal host binding (0.0.0.0) guarantees access via localhost, 127.0.0.1, and local networks, and client-side "
        "PDF generation functions without server dependencies."
    )

    # Section 9: Legal Disclaimer
    add_heading_1("9. Legal & Financial Disclaimer")
    doc.add_paragraph(
        "DealWise AI provides deterministic analytical estimates and qualitative AI interpretations for educational "
        "and decision-support purposes only. DealWise AI is not a registered financial advisor, licensed broker, or legal practice. "
        "Future property values, capital appreciation rates, and rental incomes are estimates, not guarantees. "
        "Buyers must independently verify all municipal documentation and title deeds prior to contract signing."
    )

    doc.save("DEALWISE_AI_PROJECT_REPORT.docx")
    print("Created DEALWISE_AI_PROJECT_REPORT.docx successfully.")

# ==========================================
# 2. GENERATE POWERPOINT PRESENTATION (.pptx)
# ==========================================

def create_powerpoint_presentation():
    prs = Presentation()
    # 16:9 Widescreen dimensions
    prs.slide_width = PptxInches(13.333)
    prs.slide_height = PptxInches(7.5)
    blank_layout = prs.slide_layouts[6] # blank layout

    BG_COLOR = PptxRGBColor(13, 13, 17)       # #0D0D11
    CARD_BG = PptxRGBColor(24, 24, 30)        # #18181E
    WHITE = PptxRGBColor(255, 255, 255)
    LIGHT_GRAY = PptxRGBColor(212, 212, 216)  # #D4D4D8
    MUTED_GRAY = PptxRGBColor(161, 161, 170)  # #A1A1AA
    ACCENT_LINE = PptxRGBColor(60, 60, 70)

    def set_slide_background(slide):
        bg = slide.shapes.add_shape(1, 0, 0, prs.slide_width, prs.slide_height) # msoShapeRectangle
        bg.fill.solid()
        bg.fill.fore_color.rgb = BG_COLOR
        bg.line.fill.background()

    def add_card(slide, left, top, width, height, title, body_bullets):
        card = slide.shapes.add_shape(1, left, top, width, height)
        card.fill.solid()
        card.fill.fore_color.rgb = CARD_BG
        card.line.color.rgb = ACCENT_LINE
        card.line.width = PptxPt(1)

        tf = card.text_frame
        tf.word_wrap = True
        tf.margin_left = PptxInches(0.25)
        tf.margin_right = PptxInches(0.25)
        tf.margin_top = PptxInches(0.25)
        tf.margin_bottom = PptxInches(0.25)

        p = tf.paragraphs[0]
        p.text = title
        p.font.size = PptxPt(16)
        p.font.bold = True
        p.font.color.rgb = WHITE
        p.space_after = PptxPt(10)

        for b in body_bullets:
            p2 = tf.add_paragraph()
            p2.text = f"• {b}"
            p2.font.size = PptxPt(12)
            p2.font.color.rgb = LIGHT_GRAY
            p2.space_after = PptxPt(6)

    def add_header(slide, tag_text, title_text):
        tag_box = slide.shapes.add_textbox(PptxInches(0.8), PptxInches(0.5), PptxInches(11.7), PptxInches(0.4))
        tf_t = tag_box.text_frame
        p_t = tf_t.paragraphs[0]
        p_t.text = tag_text.upper()
        p_t.font.size = PptxPt(11)
        p_t.font.color.rgb = MUTED_GRAY
        p_t.font.bold = True

        title_box = slide.shapes.add_textbox(PptxInches(0.8), PptxInches(0.85), PptxInches(11.7), PptxInches(0.8))
        tf_title = title_box.text_frame
        p_title = tf_title.paragraphs[0]
        p_title.text = title_text
        p_title.font.size = PptxPt(26)
        p_title.font.bold = True
        p_title.font.color.rgb = WHITE

    # Slide 1: Title Slide
    s1 = prs.slides.add_slide(blank_layout)
    set_slide_background(s1)

    t_box = s1.shapes.add_textbox(PptxInches(1.0), PptxInches(1.8), PptxInches(11.3), PptxInches(3.8))
    tf1 = t_box.text_frame
    p1 = tf1.paragraphs[0]
    p1.text = "DEALWISE AI"
    p1.font.size = PptxPt(52)
    p1.font.bold = True
    p1.font.color.rgb = WHITE
    p1.space_after = PptxPt(8)

    p2 = tf1.add_paragraph()
    p2.text = "Know the deal before you buy."
    p2.font.size = PptxPt(24)
    p2.font.color.rgb = MUTED_GRAY
    p2.space_after = PptxPt(18)

    p3 = tf1.add_paragraph()
    p3.text = "AI-Powered Real Estate Deal Analyzer & Investment Intelligence Platform\nAuthor: Panchaksharayya | October 2026\nRepository: https://github.com/panchaksharayya12/DealwiseAi"
    p3.font.size = PptxPt(14)
    p3.font.color.rgb = LIGHT_GRAY

    # Slide 2: The Core Problem
    s2 = prs.slides.add_slide(blank_layout)
    set_slide_background(s2)
    add_header(s2, "01. Problem Analysis", "Why Real Estate Buyers Make Costly Mistakes")
    add_card(s2, PptxInches(0.8), PptxInches(1.8), PptxInches(5.6), PptxInches(2.4), "The Price vs. Value Distortion", [
        "Buyers evaluate properties solely by location and gross asking price.",
        "Fails to analyze net capitalization yields, operating overhead, and actual returns.",
        "Multi-crore decisions made with emotional bias rather than financial diligence."
    ])
    add_card(s2, PptxInches(6.8), PptxInches(1.8), PptxInches(5.6), PptxInches(2.4), "Negative Monthly Carry", [
        "High LTV financing paired with society maintenance causes cashflow deficits.",
        "Monthly EMI and maintenance frequently exceed rental income.",
        "Forces owners to fund recurring shortfalls from active income."
    ])
    add_card(s2, PptxInches(0.8), PptxInches(4.5), PptxInches(5.6), PptxInches(2.4), "AI Financial Hallucinations", [
        "Generic LLMs fabricate loan amortization schedules and compound interest.",
        "Math is often invented, producing dangerously erroneous investment guidance.",
        "Mathematical modeling must be 100% deterministic."
    ])
    add_card(s2, PptxInches(6.8), PptxInches(4.5), PptxInches(5.6), PptxInches(2.4), "Broker-Biased Platforms", [
        "Traditional property portals earn brokerage and listing fees.",
        "Their UX is optimized to sell properties, not protect buyers from bad investments.",
        "Zero objective risk diagnostics or comparative stress tests."
    ])

    # Slide 3: The Solution
    s3 = prs.slides.add_slide(blank_layout)
    set_slide_background(s3)
    add_header(s3, "02. The DealWise Solution", "Objective Quantitative Property Diligence")
    add_card(s3, PptxInches(0.8), PptxInches(1.8), PptxInches(3.6), PptxInches(5.0), "Deterministic Math", [
        "100% verifiable TypeScript routines.",
        "Price per sq.ft, rental yields, monthly EMI, and total interest calculated with actuarial accuracy.",
        "Zero financial hallucination.",
        "Transparent equations accessible to the user."
    ])
    add_card(s3, PptxInches(4.8), PptxInches(1.8), PptxInches(3.6), PptxInches(5.0), "5-Pillar Deal Score", [
        "Weighted score from 0 to 100.",
        "Evaluates Financial Health (35%), Yield (20%), Price (20%), Debt (15%), and Risk (10%).",
        "Clear classifications: Strong Deal, Fair Deal, Needs Review, Risky Deal.",
        "Actionable diagnostic rationale."
    ])
    add_card(s3, PptxInches(8.8), PptxInches(1.8), PptxInches(3.6), PptxInches(5.0), "Contextual AI & Reports", [
        "AI strictly deployed for qualitative synthesis and seller negotiation strategies.",
        "Context-aware interactive Q&A assistant.",
        "Dynamic 5-year growth simulations.",
        "One-click executive-grade PDF due diligence reports."
    ])

    # Slide 4: Actuarial Financial Engine
    s4 = prs.slides.add_slide(blank_layout)
    set_slide_background(s4)
    add_header(s4, "03. Financial Modeling", "Mathematical & Actuarial Precision")
    add_card(s4, PptxInches(0.8), PptxInches(1.8), PptxInches(5.6), PptxInches(5.0), "Yield & Valuation Equations", [
        "Price / sq.ft = Asking Price / Built-up Area",
        "Annual Rent = Expected Monthly Rent * 12",
        "Gross Rental Yield (%) = (Annual Rent / Asking Price) * 100",
        "Annual Maintenance = Monthly Maintenance * 12",
        "Net Rental Income = Annual Rent - Annual Maintenance",
        "Net Rental Yield (%) = (Net Rental Income / Asking Price) * 100",
        "Cash-on-Cash ROI (%) = (Annual Net Cashflow / Down Payment) * 100"
    ])
    add_card(s4, PptxInches(6.8), PptxInches(1.8), PptxInches(5.6), PptxInches(5.0), "Debt Amortization & Carry", [
        "Loan Amount = max(0, Asking Price - Down Payment)",
        "Standard Reducing Balance EMI Formula:",
        "  EMI = P * r * (1+r)^n / ((1+r)^n - 1)",
        "  P = Loan Amount",
        "  r = Annual Interest Rate / (12 * 100)",
        "  n = Loan Tenure (years) * 12",
        "Total Interest = (EMI * n) - Loan Amount",
        "Monthly Cashflow Carry = Rent - Maintenance - Monthly EMI",
        "Loan-to-Value (LTV) Ratio = (Loan Amount / Asking Price) * 100"
    ])

    # Slide 5: The 5-Pillar Deal Score
    s5 = prs.slides.add_slide(blank_layout)
    set_slide_background(s5)
    add_header(s5, "04. Scoring Methodology", "Transparent Weighted Deal Score (0–100)")
    add_card(s5, PptxInches(0.8), PptxInches(1.8), PptxInches(5.6), PptxInches(5.0), "Score Composition Weights", [
        "Financial Health (35%): Net rental yield vs. cost of capital and monthly cashflow carry.",
        "Rental Yield Power (20%): Gross rental yield benchmarked against metro averages (>5.5% top tier).",
        "Price Efficiency (20%): Price per sq.ft affordability and space yield sanity.",
        "Loan & Debt Burden (15%): Loan-to-Value (LTV) exposure and monthly debt service coverage.",
        "Risk Buffer (10%): Asset age depreciation, maintenance drag ratio, and dedicated parking."
    ])
    add_card(s5, PptxInches(6.8), PptxInches(1.8), PptxInches(5.6), PptxInches(5.0), "Diagnostic Tiers", [
        "80 – 100: Strong Deal",
        "  Solid fundamentals, high yields, sustainable debt, positive cashflow.",
        "60 – 79: Fair Deal",
        "  Viable opportunity; minor price negotiation or higher down payment advised.",
        "40 – 59: Needs Review",
        "  Friction detected: yield compression or excessive monthly debt burden.",
        "0 – 39: Risky Deal",
        "  High leverage danger, severe monthly deficit, or inflated asking price."
    ])

    # Slide 6: Risk Analysis & Due Diligence Checklist
    s6 = prs.slides.add_slide(blank_layout)
    set_slide_background(s6)
    add_header(s6, "05. Risk Taxonomy", "Automated Diagnostics & Buyer Due Diligence")
    add_card(s6, PptxInches(0.8), PptxInches(1.8), PptxInches(5.6), PptxInches(5.0), "Automated Risk Flags", [
        "High LTV Leverage (>75% LTV): Heightened debt vulnerability and interest expense.",
        "Monthly Cashflow Deficit: Highlights when EMI + Maintenance exceeds monthly rent.",
        "Yield Compression (<2.5%): Identifies speculative deals reliant solely on future appreciation.",
        "Maintenance Drag (>18% of rent): Warns of recurring operational erosion.",
        "Construction Age (>15 yrs): Flags upcoming structural and society sinking fund liabilities."
    ])
    add_card(s6, PptxInches(6.8), PptxInches(1.8), PptxInches(5.6), PptxInches(5.0), "8-Point Due Diligence Checklist", [
        "1. 30-year unbroken chain title deeds & ownership documents.",
        "2. RERA project registration, sanction plans & CC/OC certificates.",
        "3. Encumbrance Certificate (EC) verifying absence of legal mortgages.",
        "4. Registered sale deeds of 500m comparable properties.",
        "5. Society sinking fund health & upcoming capital repair assessments.",
        "6. Local broker validation of realistic prevailing market rents.",
        "7. Property tax paid receipts and utility clearance certificates.",
        "8. Bank APF mortgage pre-approval from leading institutions."
    ])

    # Slide 7: 5-Year Projections & Comparison
    s7 = prs.slides.add_slide(blank_layout)
    set_slide_background(s7)
    add_header(s7, "06. Interactive Features", "Projections & Multi-Property Benchmarking")
    add_card(s7, PptxInches(0.8), PptxInches(1.8), PptxInches(5.6), PptxInches(5.0), "5-Year Growth Projections", [
        "Dynamic Recharts SVG Area Graph modeling holding horizon.",
        "Interactive Appreciation Slider: Adjust scenario from 0% to 12% annually.",
        "Models asset appreciation compounded with conservative 4% annual rental step-ups.",
        "Plots combined capital gains alongside accumulated cash distributions.",
        "Statutory regulatory labeling: Illustrative estimate, non-guaranteed."
    ])
    add_card(s7, PptxInches(6.8), PptxInches(1.8), PptxInches(5.6), PptxInches(5.0), "Multi-Property Benchmarking", [
        "Side-by-side evaluation of 2 or more candidate properties.",
        "Automated category winners: Best Yield, Lowest Price/sq.ft, Best Cashflow, Top Score.",
        "Preloaded presets: Bengaluru 2 BHK, Hyderabad 3 BHK, Goa 4 BHK Villa.",
        "Dynamic \"DealWise Recommended Deal\" banner highlighting optimal risk-adjusted return."
    ])

    # Slide 8: System Architecture & Stack
    s8 = prs.slides.add_slide(blank_layout)
    set_slide_background(s8)
    add_header(s8, "07. Engineering", "System Architecture & Production Stack")
    add_card(s8, PptxInches(0.8), PptxInches(1.8), PptxInches(3.6), PptxInches(5.0), "Frontend Client", [
        "React 18 & TypeScript",
        "Vite Build Tooling",
        "Tailwind CSS & Liquid Glass",
        "Recharts SVG Analytics",
        "Lucide React Iconography",
        "jsPDF Vector Report Engine"
    ])
    add_card(s8, PptxInches(4.8), PptxInches(1.8), PptxInches(3.6), PptxInches(5.0), "Backend REST API", [
        "Node.js & Express (TypeScript)",
        "Deterministic Valuation Services",
        "OpenAI GPT-4o-mini Orchestration",
        "Automatic Local Fallback Engine",
        "Universal 0.0.0.0 Host Binding",
        "Robust Error & Exception Guards"
    ])
    add_card(s8, PptxInches(8.8), PptxInches(1.8), PptxInches(3.6), PptxInches(5.0), "Persistence & Cloud", [
        "Supabase PostgreSQL Integration",
        "Offline Browser localStorage Sync",
        "Zero-Config Local Development",
        "Production-ready Docker / Cloud Deployable"
    ])

    # Slide 9: Conclusion
    s9 = prs.slides.add_slide(blank_layout)
    set_slide_background(s9)
    add_header(s9, "08. Summary", "DealWise AI — Deliverables & Impact")
    add_card(s9, PptxInches(1.8), PptxInches(1.8), PptxInches(9.7), PptxInches(4.8), "Key Takeaways", [
        "Production-Grade Platform: 100% functional, responsive, and rigorously verified.",
        "Mathematical Rigor: Zero AI hallucination on financial calculations; full transparency.",
        "Premium VEX Visuals: Liquid glass styling, cinematic hero video, and zero gradients.",
        "Comprehensive Deliverables: Full-stack codebase, Word Document Report, PowerPoint Deck, and Live App.",
        "Live Application URL: http://localhost:5173",
        "GitHub Repository: https://github.com/panchaksharayya12/DealwiseAi"
    ])

    prs.save("DEALWISE_AI_PRESENTATION.pptx")
    print("Created DEALWISE_AI_PRESENTATION.pptx successfully.")

if __name__ == "__main__":
    create_word_report()
    create_powerpoint_presentation()
